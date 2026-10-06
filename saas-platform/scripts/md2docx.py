#!/usr/bin/env python3
"""Convert the Priya project specification Markdown into a formatted .docx.

Deliberately dependency-light: python-docx only. Handles headings, paragraphs
with **bold** / `code` / *italic*, GFM pipe tables, fenced code blocks (including
mermaid, which Word cannot render and is therefore kept as source text),
blockquotes, horizontal rules, ordered lists, and task-list checkboxes.
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

SRC = sys.argv[1]
DST = sys.argv[2]

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")
CHECKBOX = re.compile(r"^\[([ xX])\]\s*")
ORDERED = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
UNORDERED = re.compile(r"^(\s*)[-*+]\s+(.*)$")
SEPARATOR = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
FENCE = re.compile(r"^\s*```")


def shade(element, hex_fill):
    tc_pr = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_inline(paragraph, text, base_size=None):
    """Render **bold**, `code` and *italic* spans into runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((base_size or 10.5) - 0.5)
            run.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if base_size:
            run.font.size = Pt(base_size)
    return paragraph


def flush_code(doc, lines):
    """A fenced block becomes monospace, shaded, line-per-paragraph."""
    for line in lines or [""]:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.left_indent = Inches(0.18)
        run = para.add_run(line if line.strip() else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        shade(para._p, "F4F4F4")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def flush_table(doc, rows):
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for idx, cells in enumerate(rows):
        row = table.add_row()
        for col in range(width):
            cell = row.cells[col]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            text = cells[col] if col < len(cells) else ""
            add_inline(para, text, base_size=9)
            if idx == 0:
                for run in para.runs:
                    run.bold = True
                shade(cell._tc, "EFEFEF")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    with open(SRC, encoding="utf-8") as handle:
        lines = handle.read().split("\n")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)

    in_code = False
    code_lines: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if FENCE.match(line):
            if in_code:
                flush_code(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        stripped = line.strip()

        if stripped.startswith("|") and i + 1 < len(lines) and SEPARATOR.match(lines[i + 1].strip()):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            flush_table(doc, rows)
            continue

        if not stripped:
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(8)
            run = para.add_run("─" * 46)
            run.font.color.rgb = RGBColor(0xC0, 0xC0, 0xC0)
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).replace("**", "").replace("`", "")
            para = doc.add_heading(level=min(level, 4))
            run = para.add_run(text)
            if level == 1:
                run.font.size = Pt(20)
            elif level == 2:
                run.font.size = Pt(15)
            else:
                run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            i += 1
            continue

        if stripped.startswith("> "):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.3)
            add_inline(para, stripped[2:], base_size=10)
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            i += 1
            continue

        task = UNORDERED.match(line)
        if task:
            indent = len(task.group(1))
            body = task.group(2)
            mark = CHECKBOX.match(body)
            style = "List Bullet" if indent < 2 else "List Bullet 2"
            para = doc.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(3)
            if mark:
                body = body[mark.end():]
                box = "☒  " if mark.group(1).lower() == "x" else "☐  "
                para.add_run(box)
            add_inline(para, body)
            i += 1
            continue

        num = ORDERED.match(line)
        if num:
            para = doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(3)
            add_inline(para, num.group(3))
            i += 1
            continue

        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(6)
        add_inline(para, stripped)
        i += 1

    doc.save(DST)
    print(f"wrote {DST}")


if __name__ == "__main__":
    main()
